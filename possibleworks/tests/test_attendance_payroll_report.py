# Copyright (c) 2026, Possibleworks and contributors
# For license information, please see license.txt

"""Tests for possibleworks.hr_documents.attendance_payroll_report: the payroll
attendance-exception report (Absent/Half-Day rows only, a holiday flag that
stays independent of Attendance.status, pending-regularization detection, and
on-shift vs. off-shift check-in enrichment).

NOTE ON FIXTURES: this module builds Company/Employee/Holiday List/Payroll
Period fixtures directly with fixed names rather than relying on
IntegrationTestCase's recursive test-record auto-creation (that mechanism --
and the IGNORE_TEST_RECORD_DEPENDENCIES override used by
hr_documents/doctype/form_16/test_form_16.py -- only applies to test modules
that live inside a doctype's own folder; this module lives in tests/, so it
manages its own fixed-name fixtures and cleans them up with
delete_doc_if_exists instead).

NOTE ON ISOLATION: creating an Employee fires the Observer, and Employee is in
IMMEDIATE_SEND_DOCTYPES, whose branch calls frappe.db.commit()
(possibleworks/observer/observer.py). Every test that creates an Employee
therefore neutralises WorkflowEventObserver.should_process, exactly as
test_employee.py does.
"""

from unittest.mock import patch

import frappe
from frappe.tests import IntegrationTestCase
from frappe.utils import getdate

import erpnext
from erpnext.setup.doctype.employee.test_employee import make_employee

from hrms.payroll.doctype.salary_slip.test_salary_slip import make_holiday_list
from hrms.tests.test_utils import add_date_to_holiday_list, create_company

from possibleworks.hr_documents.attendance_payroll_report.attendance_payroll_report import (
	_build_report_rows,
	_parse_emails,
	daily_attendance_report_dispatch,
	send_attendance_exception_report,
)
from possibleworks.hr_documents.attendance_payroll_report.excel_builder import (
	HOLIDAY_FILL,
	build_attendance_exception_workbook,
)
from possibleworks.observer.observer import WorkflowEventObserver
from possibleworks.tests.site_fixtures import site_mandatory_values

COMPANY = erpnext.get_default_company() or "_Test Company"
HOLIDAY_LIST = "_Test Attendance Report Holidays"
HOLIDAY_DATE = "2026-04-14"
PAYROLL_PERIOD = "_Test Attendance Report Period"
FROM_DATE = getdate("2026-04-01")
TO_DATE = getdate("2026-04-30")


def _make_report_employee(user, **kwargs):
	"""A test Employee that satisfies whatever THIS site marks as mandatory --
	same reasoning as hr_documents/doctype/form_16/test_form_16.py's
	make_form16_employee: a site can add its own reqd fields (e.g. griet-local
	requires employee_number) via Custom Field/Property Setter, so read the
	live meta instead of hardcoding."""
	meta = frappe.get_meta("Employee")
	extra = site_mandatory_values("Employee", exclude=("employee_number", "reports_to"))

	employee_number = meta.get_field("employee_number")
	if employee_number and employee_number.reqd:
		extra["employee_number"] = f"ATTREP-{frappe.generate_hash(length=8).upper()}"

	reports_to = meta.get_field("reports_to")
	if reports_to and reports_to.reqd:
		manager = frappe.get_all(
			"Employee",
			filters=[["status", "=", "Active"], ["user_id", "is", "set"], ["user_id", "!=", user]],
			pluck="name",
			limit=1,
		)
		if manager:
			extra["reports_to"] = manager[0]

	extra.update(kwargs)
	return make_employee(user, **extra)


def _make_attendance(employee, date, status, working_hours=None):
	doc = frappe.get_doc(
		{
			"doctype": "Attendance",
			"employee": employee,
			"attendance_date": date,
			"status": status,
			"company": COMPANY,
			"working_hours": working_hours,
		}
	)
	doc.insert()
	doc.submit()
	return doc


def _make_checkin(employee, time, log_type="IN"):
	return frappe.get_doc(
		{
			"doctype": "Employee Checkin",
			"employee": employee,
			"time": time,
			"log_type": log_type,
			"device_id": "test-device",
		}
	).insert()


def _make_draft_attendance_request(employee, from_date, to_date, reason="On Duty"):
	doc = frappe.get_doc(
		{
			"doctype": "Attendance Request",
			"employee": employee,
			"from_date": from_date,
			"to_date": to_date,
			"reason": reason,
			"company": COMPANY,
		}
	)
	doc.insert()
	return doc


def _make_open_leave_application(employee, from_date, to_date, leave_type="Leave Without Pay"):
	# LWP skips the leave-balance check (is_lwp in validate_balance_leaves), so
	# this needs no Leave Allocation fixture -- "Leave Without Pay" is a
	# standard leave type hrms seeds on install (hrms/setup.py).
	doc = frappe.get_doc(
		{
			"doctype": "Leave Application",
			"employee": employee,
			"leave_type": leave_type,
			"from_date": from_date,
			"to_date": to_date,
			"company": COMPANY,
			"status": "Open",
			"leave_approver": "Administrator",
		}
	)
	doc.insert(ignore_permissions=True)
	return doc


class TestAttendancePayrollReport(IntegrationTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()

		cls._observer_patch = patch.object(WorkflowEventObserver, "should_process", return_value=False)
		cls._observer_patch.start()
		cls.addClassCleanup(cls._observer_patch.stop)

		frappe.set_user("Administrator")

		for doctype in (
			"Attendance",
			"Employee Checkin",
			"Attendance Request",
			"Leave Application",
			"Holiday List Assignment",
		):
			frappe.db.delete(doctype)

		cls.holiday_list = make_holiday_list(
			list_name=HOLIDAY_LIST, from_date=FROM_DATE, to_date=TO_DATE, add_weekly_offs=False
		)
		add_date_to_holiday_list(HOLIDAY_DATE, cls.holiday_list)

		cls.employee = _make_report_employee("attendance.report@example.com", company=COMPANY)

		hla = frappe.new_doc("Holiday List Assignment")
		hla.applicable_for = "Employee"
		hla.assigned_to = cls.employee
		hla.holiday_list = cls.holiday_list
		hla.employee_company = COMPANY
		hla.from_date = FROM_DATE
		hla.insert()
		hla.submit()

		frappe.delete_doc_if_exists("Payroll Period", PAYROLL_PERIOD, force=True)
		cls.payroll_period = frappe.get_doc(
			{
				"doctype": "Payroll Period",
				"name": PAYROLL_PERIOD,
				"company": COMPANY,
				"start_date": FROM_DATE,
				"end_date": TO_DATE,
				"custom_payroll_type": "Monthly",
				"custom_to_emails": "payroll@example.com, payroll-lead@example.com",
				"custom_cc_emails": "hr@example.com",
			}
		).insert()

	def setUp(self):
		frappe.set_user("Administrator")
		# Unscoped, not just self.employee: a couple of tests create their own
		# one-off employee (e.g. the missing-holiday-list-coverage case), and
		# everything here lives inside the class-level rollback anyway.
		for doctype in ("Attendance", "Employee Checkin", "Attendance Request", "Leave Application"):
			frappe.db.delete(doctype)

	def tearDown(self):
		frappe.set_user("Administrator")

	# -------------------------------------------------------------------
	# _parse_emails
	# -------------------------------------------------------------------

	def test_parse_emails_splits_dedupes_and_drops_invalid_tokens(self):
		value = "a@example.com, b@example.com\nc@example.com; a@example.com, not-an-email, "
		self.assertEqual(
			_parse_emails(value), ["a@example.com", "b@example.com", "c@example.com"]
		)

	def test_parse_emails_handles_empty_input(self):
		self.assertEqual(_parse_emails(None), [])
		self.assertEqual(_parse_emails(""), [])

	# -------------------------------------------------------------------
	# _build_report_rows
	# -------------------------------------------------------------------

	def test_only_absent_and_half_day_rows_are_included(self):
		_make_attendance(self.employee, "2026-04-05", "Present")
		_make_attendance(self.employee, "2026-04-06", "Absent")
		_make_attendance(self.employee, "2026-04-07", "Half Day")

		rows = _build_report_rows(COMPANY, FROM_DATE, TO_DATE)

		statuses = {(r["date"], r["status"]) for r in rows}
		self.assertEqual(statuses, {("2026-04-06", "Absent"), ("2026-04-07", "Half Day")})

	def test_holiday_is_flagged_even_when_status_is_absent(self):
		"""The exact gap this feature closes: a holiday must show up even when
		Attendance already says Absent -- Status is never overwritten."""
		_make_attendance(self.employee, HOLIDAY_DATE, "Absent")

		rows = _build_report_rows(COMPANY, FROM_DATE, TO_DATE)

		self.assertEqual(len(rows), 1)
		row = rows[0]
		self.assertEqual(row["status"], "Absent")
		self.assertTrue(row["is_holiday"])
		self.assertTrue(row["holiday"])

	def test_non_holiday_absence_has_no_holiday_flag(self):
		_make_attendance(self.employee, "2026-04-06", "Absent")

		rows = _build_report_rows(COMPANY, FROM_DATE, TO_DATE)

		self.assertEqual(len(rows), 1)
		self.assertFalse(rows[0]["is_holiday"])
		self.assertEqual(rows[0]["holiday"], "")

	def test_working_hours_is_pulled_from_attendance_doc(self):
		_make_attendance(self.employee, "2026-04-06", "Half Day", working_hours=4.5)

		rows = _build_report_rows(COMPANY, FROM_DATE, TO_DATE)

		self.assertEqual(len(rows), 1)
		self.assertEqual(rows[0]["working_hours"], 4.5)

	def test_employee_with_no_holiday_list_coverage_does_not_break_the_report(self):
		"""possibleworks.utils.holiday_utils.get_holidays_for_employee throws when an
		employee has no Holiday List Assignment/default at all -- one employee's
		missing HR setup must not take down the whole report."""
		other_employee = _make_report_employee("attendance.report.no-holiday-list@example.com", company=COMPANY)
		_make_attendance(other_employee, "2026-04-06", "Absent")

		rows = _build_report_rows(COMPANY, FROM_DATE, TO_DATE)

		self.assertEqual(len(rows), 1)
		self.assertFalse(rows[0]["is_holiday"])

	def test_draft_attendance_request_shows_as_pending(self):
		_make_attendance(self.employee, "2026-04-08", "Half Day")
		_make_draft_attendance_request(self.employee, "2026-04-08", "2026-04-08")

		rows = _build_report_rows(COMPANY, FROM_DATE, TO_DATE)

		self.assertEqual(len(rows), 1)
		self.assertIn("Attendance Request (Draft)", rows[0]["pending_request"])

	def test_submitted_attendance_request_updates_attendance_and_drops_out_of_exceptions(self):
		"""Submitting an Attendance Request moves Attendance off Absent/Half Day
		directly (create_or_update_attendance) -- so once actioned, it is no
		longer "pending" and no longer an exception row at all."""
		_make_attendance(self.employee, "2026-04-09", "Absent")
		request = _make_draft_attendance_request(self.employee, "2026-04-09", "2026-04-09", reason="On Duty")
		request.submit()

		rows = _build_report_rows(COMPANY, FROM_DATE, TO_DATE)
		self.assertEqual([r for r in rows if r["date"] == "2026-04-09"], [])

	def test_open_leave_application_shows_as_pending(self):
		_make_attendance(self.employee, "2026-04-10", "Half Day")
		_make_open_leave_application(self.employee, "2026-04-10", "2026-04-10")

		rows = _build_report_rows(COMPANY, FROM_DATE, TO_DATE)

		self.assertEqual(len(rows), 1)
		self.assertIn("Leave Application (Open)", rows[0]["pending_request"])

	def test_checkins_split_between_normal_and_offshift(self):
		_make_attendance(self.employee, "2026-04-11", "Absent")
		on_shift = _make_checkin(self.employee, "2026-04-11 09:10:00", "IN")
		frappe.db.set_value("Employee Checkin", on_shift.name, "offshift", 0)
		_make_checkin(self.employee, "2026-04-11 21:30:00", "IN")  # naturally offshift=1, no Shift Assignment

		rows = _build_report_rows(COMPANY, FROM_DATE, TO_DATE)

		self.assertEqual(len(rows), 1)
		self.assertIn("09:10", rows[0]["checkins"])
		self.assertIn("21:30", rows[0]["offshift_checkins"])
		self.assertNotIn("21:30", rows[0]["checkins"])

	def test_no_exception_rows_returns_empty_list(self):
		self.assertEqual(_build_report_rows(COMPANY, FROM_DATE, TO_DATE), [])

	# -------------------------------------------------------------------
	# excel_builder
	# -------------------------------------------------------------------

	def test_workbook_colors_only_the_date_cell_and_only_on_holidays(self):
		from openpyxl import load_workbook
		from io import BytesIO

		rows = [
			{"employee": "E1", "date": "2026-04-06", "status": "Absent", "is_holiday": False},
			{"employee": "E2", "date": HOLIDAY_DATE, "status": "Absent", "is_holiday": True, "holiday": "Test Holiday"},
		]
		xlsx_bytes = build_attendance_exception_workbook(rows)
		wb = load_workbook(BytesIO(xlsx_bytes))
		ws = wb.active

		header = [c.value for c in ws[1]]
		date_col = header.index("Date") + 1
		status_col = header.index("Status") + 1

		non_holiday_date_cell = ws.cell(row=2, column=date_col)
		holiday_date_cell = ws.cell(row=3, column=date_col)
		status_cell = ws.cell(row=3, column=status_col)

		self.assertIsNone(non_holiday_date_cell.fill.fill_type)
		self.assertEqual(holiday_date_cell.fill.fill_type, "solid")
		self.assertEqual(holiday_date_cell.fill.fgColor.rgb, HOLIDAY_FILL.fgColor.rgb)
		# Status is never colored -- the holiday fill lives only on the Date cell.
		self.assertIsNone(status_cell.fill.fill_type)

	# -------------------------------------------------------------------
	# permissions
	# -------------------------------------------------------------------

	def test_unauthorized_role_is_rejected(self):
		user = "attendance.report.viewer@example.com"
		if not frappe.db.exists("User", user):
			frappe.get_doc(
				{
					"doctype": "User",
					"email": user,
					"first_name": "Report Viewer",
					"send_welcome_email": 0,
					"roles": [{"doctype": "Has Role", "role": "Employee"}],
				}
			).insert(ignore_permissions=True)

		frappe.set_user(user)
		try:
			self.assertRaises(
				frappe.PermissionError,
				send_attendance_exception_report,
				PAYROLL_PERIOD,
				str(FROM_DATE),
				str(TO_DATE),
			)
		finally:
			frappe.set_user("Administrator")

	# -------------------------------------------------------------------
	# send_attendance_exception_report (end to end)
	# -------------------------------------------------------------------

	@patch("possibleworks.hr_documents.attendance_payroll_report.attendance_payroll_report.frappe.sendmail")
	def test_send_report_emails_xlsx_and_updates_tracking_field(self, mock_sendmail):
		_make_attendance(self.employee, "2026-04-06", "Absent")

		result = send_attendance_exception_report(PAYROLL_PERIOD, str(FROM_DATE), str(TO_DATE))

		self.assertTrue(result["sent"])
		self.assertEqual(result["rows"], 1)

		mock_sendmail.assert_called_once()
		call_kwargs = mock_sendmail.call_args.kwargs
		self.assertEqual(call_kwargs["recipients"], ["payroll@example.com", "payroll-lead@example.com"])
		self.assertEqual(call_kwargs["cc"], ["hr@example.com"])
		self.assertEqual(len(call_kwargs["attachments"]), 1)
		self.assertTrue(call_kwargs["attachments"][0]["fname"].endswith(".xlsx"))

		self.assertEqual(
			getdate(frappe.db.get_value("Payroll Period", PAYROLL_PERIOD, "custom_last_auto_report_sent_till")),
			TO_DATE,
		)

	def test_send_report_without_recipients_throws(self):
		# Its own company, purely so this Payroll Period's date range can't
		# overlap-validate against the shared class-level one above.
		company = create_company("_Test Attendance Report No Recipients Co").name

		frappe.delete_doc_if_exists("Payroll Period", "_Test No Recipients Period", force=True)
		period = frappe.get_doc(
			{
				"doctype": "Payroll Period",
				"name": "_Test No Recipients Period",
				"company": company,
				"start_date": FROM_DATE,
				"end_date": TO_DATE,
				"custom_payroll_type": "Monthly",
			}
		).insert()

		self.assertRaises(
			frappe.ValidationError,
			send_attendance_exception_report,
			period.name,
			str(FROM_DATE),
			str(TO_DATE),
		)

	# -------------------------------------------------------------------
	# daily_attendance_report_dispatch
	# -------------------------------------------------------------------

	@patch("possibleworks.hr_documents.attendance_payroll_report.attendance_payroll_report.send_attendance_exception_report")
	@patch("possibleworks.hr_documents.attendance_payroll_report.attendance_payroll_report.nowdate")
	def test_dispatch_sends_for_a_just_completed_cycle_then_skips_it_next_run(self, mock_nowdate, mock_send):
		# "Today" = 2026-05-01, so "yesterday" = 2026-04-30 = the last day of a
		# Monthly cycle -- deterministic regardless of when this test actually runs.
		mock_nowdate.return_value = "2026-05-01"

		# Its own company: this period spans the whole year, which would
		# overlap-validate against the shared Apr-2026 class-level period above
		# if it were on the same company.
		company = create_company("_Test Attendance Report Dispatch Co").name

		year_period = "_Test Attendance Report Year"
		frappe.delete_doc_if_exists("Payroll Period", year_period, force=True)
		frappe.get_doc(
			{
				"doctype": "Payroll Period",
				"name": year_period,
				"company": company,
				"start_date": "2026-01-01",
				"end_date": "2026-12-31",
				"custom_payroll_type": "Monthly",
			}
		).insert()

		daily_attendance_report_dispatch()
		mock_send.assert_called_once_with(year_period, getdate("2026-04-01"), getdate("2026-04-30"))

		mock_send.reset_mock()
		frappe.db.set_value("Payroll Period", year_period, "custom_last_auto_report_sent_till", "2026-04-30")

		daily_attendance_report_dispatch()
		mock_send.assert_not_called()
